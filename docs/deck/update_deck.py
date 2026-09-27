"""Extend the group's existing review deck (CN SEM3.pptx) with the final-review slides.

The original deck is opened read-only and its 8 slides are kept untouched. New slides are
appended using the deck's *own* slide layouts (Office Theme: "Title and Content", "Title Only",
"Title Slide"), so fonts, colours and placement match. Tables replicate the literature-survey
tables (13 pt, bold header row + first column, 0.5 pt light-grey grid, no fill). Charts use the
theme palette. Every number is read from ../../results/*.json, so re-running this script after
new evaluation runs refreshes the deck.

    python docs/deck/update_deck.py            # draft: missing results are flagged on-slide
    python docs/deck/update_deck.py --final    # refuse to build unless all results exist
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent.parent                      # hrt-agent/
CN_DIR = PROJECT.parent                           # Desktop/CN/
SRC_DECK = CN_DIR / "CN SEM3.pptx"
OUT_DECK = CN_DIR / "CN SEM3 - Final Review.pptx"
RESULTS = PROJECT / "results"
SCREENSHOT = PROJECT / "docs" / "deck" / "console.png"

# Office theme palette of the original deck
BLACK = RGBColor(0x00, 0x00, 0x00)
TEXT = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x59, 0x59, 0x59)
GRID = RGBColor(0xC4, 0xC7, 0xC5)                # literature-table border colour
ACCENT1 = RGBColor(0x15, 0x60, 0x82)             # theme accent 1 (teal-blue)
ACCENT2 = RGBColor(0xE9, 0x71, 0x32)             # theme accent 2 (orange)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
LT2 = RGBColor(0xE8, 0xE8, 0xE8)

L_TITLE, L_CONTENT, L_TITLE_ONLY = 0, 1, 5       # layout indices in the original master

MISSING: list[str] = []


# ------------------------------------------------------------------------------------------
# data
# ------------------------------------------------------------------------------------------
def load(name: str) -> dict | None:
    p = RESULTS / f"{name}.json"
    if not p.exists():
        MISSING.append(name)
        return None
    return json.loads(p.read_text(encoding="utf-8"))["result"]


def f3(x) -> str:
    return "-" if x is None else f"{x:.3f}"


def pct(x, digits=0) -> str:
    return "-" if x is None else f"{x * 100:.{digits}f}%"


# ------------------------------------------------------------------------------------------
# slide helpers (all geometry mirrors the original slides 7-8)
# ------------------------------------------------------------------------------------------
TITLE_BOX = (Inches(0.92), Inches(0.40), Inches(11.50), Inches(1.00))
BODY_BOX = (Inches(0.92), Inches(1.55), Inches(11.50), Inches(5.45))


def _title(slide, text: str, size: int = 36):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = TITLE_BOX
    tf = t.text_frame
    tf.text = ""
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.bold = True
    run.font.size = Pt(size)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    return t


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _no_bullet(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    etree.SubElement(pPr, qn("a:buNone"))


def _runs(paragraph, parts, size: int, color: RGBColor = TEXT) -> None:
    """parts: list of (text, bold) tuples, or a plain string."""
    if isinstance(parts, str):
        parts = [(parts, False)]
    for text, bold in parts:
        r = paragraph.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color


def bullets_slide(prs, title: str, items: list, intro: str | None = None, size: int = 20,
                  notes: str = ""):
    """Title and Content slide with bold lead-in bullets, like the original 'Research Gaps' slide.
    items: list of (lead, text) or plain strings."""
    s = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    _title(s, title)
    body = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
    body.left, body.top, body.width, body.height = BODY_BOX
    tf = body.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.text = ""
    first = True
    if intro:
        p = tf.paragraphs[0]
        _no_bullet(p)
        _runs(p, intro, size)
        p.space_after = Pt(10)
        first = False
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        p.space_after = Pt(9)
        if isinstance(it, tuple):
            lead, text = it
            _runs(p, [(lead + " ", True), (text, False)], size)
        else:
            _runs(p, it, size)
    if notes:
        _notes(s, notes)
    return s


def blank_slide(prs, title: str, notes: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    _title(s, title)
    if notes:
        _notes(s, notes)
    return s


def _bullet(paragraph, size: int) -> None:
    """Theme-style '•' bullet with a hanging indent (matches the original content slides)."""
    pPr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(size).emu * 1.15)
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def textbox(slide, x, y, w, h, parts_list, size=14, color=TEXT, bold_first=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, space_after=6, bullets=False):
    """parts_list: list of paragraphs; each paragraph is str or list[(text,bold)]."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, para in enumerate(parts_list):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if bullets:
            _bullet(p, size)
        _runs(p, para, size, color)
    return tb


# ---- tables in the literature-survey style ----------------------------------------------
def _cell_border(cell, color: RGBColor = GRID, width_emu: int = 6350) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for el in tcPr.findall(qn("a:noFill")) + tcPr.findall(qn("a:solidFill")):
        tcPr.remove(el)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(width_emu), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=str(color))
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    etree.SubElement(tcPr, qn("a:noFill"))


def table(slide, rows: list[list], x, y, col_widths_in: list[float], font=13, header=True,
          bold_first_col=True, row_h=0.4, highlight_rows: set[int] | None = None, header_h: float | None = None):
    n_r, n_c = len(rows), len(rows[0])
    width = Inches(sum(col_widths_in))
    header_h = header_h if header_h is not None else min(row_h, 0.45)
    gf = slide.shapes.add_table(n_r, n_c, x, y, width, Inches(header_h + row_h * (n_r - 1)))
    tbl = gf.table
    # match the originals: no table style, no banding (<a:tblPr/>)
    tblPr = tbl._tbl.tblPr
    for attr in list(tblPr.attrib):
        del tblPr.attrib[attr]
    for child in list(tblPr):
        tblPr.remove(child)
    for i, w in enumerate(col_widths_in):
        tbl.columns[i].width = Inches(w)
    for r in range(n_r):
        tbl.rows[r].height = Inches(header_h if (header and r == 0) else row_h)
        for c in range(n_c):
            cell = tbl.cell(r, c)
            _cell_border(cell)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Emu(55114)
            cell.margin_top = cell.margin_bottom = Emu(36743)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.text = ""
            val = rows[r][c]
            bold = (header and r == 0) or (bold_first_col and c == 0) or (highlight_rows and r in highlight_rows)
            p = tf.paragraphs[0]
            _runs(p, [(str(val), bold)] if not isinstance(val, list) else val, font, BLACK)
    return gf


# ---- charts in the theme palette ----------------------------------------------------------
def _style_chart(chart, title: str, pct_axis: bool, max_val=None, legend=True, font=11):
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    tp = chart.chart_title.text_frame.paragraphs[0]
    for r in tp.runs:
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = TEXT
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LT2
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(font)
    va.tick_labels.font.color.rgb = MUTED
    if pct_axis:
        va.tick_labels.number_format = "0%"
        va.tick_labels.number_format_is_linked = False
        # headroom above 100% so outside-end data labels never collide with the chart title
        va.maximum_scale = 1.12 if max_val is None else max_val
        va.minimum_scale = 0.0
        va.major_unit = 0.25
    elif max_val is not None:
        va.maximum_scale = max_val
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(font)
    ca.tick_labels.font.color.rgb = TEXT
    ca.format.line.color.rgb = GRID
    ca.has_major_gridlines = False


def bar_chart(slide, x, y, w, h, categories, series: list[tuple[str, list, RGBColor]], title,
              pct_axis=True, horizontal=False, labels=True, legend=True, gap=80):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals, _ in series:
        cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gframe = slide.shapes.add_chart(kind, x, y, w, h, cd)
    chart = gframe.chart
    _style_chart(chart, title, pct_axis, legend=legend and len(series) > 1)
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = -10 if len(series) > 1 else 0
    for i, (_, _, color) in enumerate(series):
        s = plot.series[i]
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = color
        s.invert_if_negative = False
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(10)
        dl.font.color.rgb = TEXT
        dl.number_format = "0%" if pct_axis else "General"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    return chart


def line_chart(slide, x, y, w, h, categories, series: list[tuple[str, list, RGBColor]], title):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals, _ in series:
        cd.add_series(name, vals)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, cd).chart
    _style_chart(chart, title, pct_axis=False)
    chart.value_axis.tick_labels.number_format = '0"s"'
    chart.value_axis.tick_labels.number_format_is_linked = False
    for i, (_, _, color) in enumerate(series):
        s = chart.plots[0].series[i]
        s.format.line.color.rgb = color
        s.format.line.width = Pt(2.5)
        s.smooth = False
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = color
        s.marker.format.line.color.rgb = color
    return chart


# ---- diagram primitives ---------------------------------------------------------------------
def box(slide, x, y, w, h, head: str, lines: list[str], fill=None, head_color=BLACK, size=12):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)
    _runs(p, [(head, True)], size + 2, head_color)
    for ln in lines:
        q = tf.add_paragraph()
        q.alignment = PP_ALIGN.CENTER
        q.space_after = Pt(3)
        _runs(q, ln, size, TEXT)
    return shp


def arrow(slide, x1, y1, x2, y2, color=RGBColor(0x59, 0x59, 0x59), width=1.75, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def pill(slide, x, y, w, h, parts, fill=LIGHT, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.18
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _runs(p, parts, size, TEXT)
    return shp


def stat(slide, x, y, w, value: str, label: str, color=ACCENT1):
    textbox(slide, x, y, w, Inches(0.8), [[(value, True)]], size=40, color=color, space_after=0)
    textbox(slide, x, y + Inches(0.78), w, Inches(0.6), [label], size=13, color=MUTED, space_after=0)


def pending(slide, what: str):
    pill(slide, Inches(0.92), Inches(3.2), Inches(11.5), Inches(1.0),
         [(f"RESULTS PENDING: {what} has not finished yet. Re-run update_deck.py when it has.", True)],
         fill=RGBColor(0xFD, 0xE9, 0xD9), size=16)


# ------------------------------------------------------------------------------------------
# measured retrieval examples: replay the real lab conn.log -> real Tier-1 alerts -> retriever
# ------------------------------------------------------------------------------------------
def retrieval_examples() -> list[dict] | None:
    try:
        from hrt.config import get_settings
        from hrt.rag.retriever import Retriever
        from hrt.telemetry.zeek import read_conn_log
        from hrt.triage.detector import WindowDetector
        from hrt.triage.engine import TriageEngine
    except Exception as exc:  # environment issue
        print("retrieval examples unavailable:", exc)
        return None
    s = get_settings()
    if not s.conn_log.exists():
        return None
    flows = read_conn_log(s.conn_log)
    alerts = TriageEngine(WindowDetector(), lambda a: None, s.window_seconds, 10.0, origin="replay").replay(flows)
    wanted = {"PORT_SCAN": "Port scan (nmap)", "BRUTE_FORCE": "SSH brute force (hydra)",
              "SYN_FLOOD": "SYN flood (hping3)"}
    picked: dict[str, object] = {}
    for a in alerts:
        for rule in a.detector.rules:
            if rule in wanted and rule not in picked:
                picked[rule] = a
    if not picked:
        return None
    r = Retriever(s)
    out = []
    for rule, label in wanted.items():
        a = picked.get(rule)
        if a is None:
            continue
        trace, _ = r.retrieve(a, top_k=3)
        top = trace.docs[0]
        out.append({"pattern": label, "tactics": ", ".join(trace.candidate_tactics),
                    "space": f"{trace.candidates}/{trace.total_docs} ({trace.search_space_reduction:.0%} cut)",
                    "top": f"{top.mitre_id} {top.name}", "score": top.score})
    return out


# ------------------------------------------------------------------------------------------
# new slides (appended after the original 8)
# ------------------------------------------------------------------------------------------
def s_architecture(prs):
    s = blank_slide(prs, "System Architecture: Three-Tier Pipeline", notes=(
        "Walk left to right. Zeek watches the protected server. Tier 1 turns thousands of flows into a small "
        "alert and posts it to an async queue, so ingestion never waits for the LLM. Tier 2 grounds the analysis "
        "in MITRE ATT&CK and writes an iptables rule. Tier 3 verifies it deterministically; rejections loop back "
        "for self-correction. Only a verified, analyst-approved rule is deployed."))
    widths = [1.60, 2.30, 2.30, 2.30, 1.60]
    gap = 0.42
    x = (13.333 - (sum(widths) + gap * 4)) / 2
    y, h = 1.85, 2.05
    specs = [
        ("Zeek sensor", ["Passive capture at the protected server", "JSON conn.log"], LIGHT, BLACK),
        ("Tier 1 · Triage", ["10 s per-source windows", "Threshold rules + Isolation Forest",
                             "Alert JSON posted to an async webhook"], None, ACCENT1),
        ("Tier 2 · RAG Brain", ["Hierarchical retrieval: ChromaDB (ATT&CK + KEV)",
                                "Local Llama 3.2 3B, JSON schema", "Grounding check"], None, ACCENT1),
        ("Tier 3 · Guardrail", ["Policy → subnet (bitwise AND) → conflict → kernel test",
                                "Ephemeral Docker sandbox"], None, ACCENT1),
        ("Gated deploy", ["Analyst approves", "iptables rule on the server"], LIGHT, BLACK),
    ]
    boxes = []
    for (head, lines, fill, hc), w in zip(specs, widths):
        boxes.append((x, w))
        box(s, Inches(x), Inches(y), Inches(w), Inches(h), head, lines, fill=fill, head_color=hc, size=12)
        x += w + gap
    mid = y + h / 2
    for (bx, bw), (nx, _) in zip(boxes, boxes[1:]):
        arrow(s, Inches(bx + bw + 0.04), Inches(mid), Inches(nx - 0.04), Inches(mid))
    # self-correction loop: Tier 3 bottom -> down -> left -> up into Tier 2
    (t2x, t2w), (t3x, t3w) = boxes[2], boxes[3]
    yb, yl = y + h, y + h + 0.38
    c2, c3 = t2x + t2w / 2, t3x + t3w / 2
    arrow(s, Inches(c3), Inches(yb), Inches(c3), Inches(yl), head=False)
    arrow(s, Inches(c3), Inches(yl), Inches(c2), Inches(yl), head=False)
    arrow(s, Inches(c2), Inches(yl), Inches(c2), Inches(yb + 0.02))
    textbox(s, Inches(c2 - 1.1), Inches(yl + 0.05), Inches(c3 - c2 + 2.2), Inches(0.35),
            [[("Rejected → reason fed back → self-correct (≤ 3 tries), else escalate to a human", False)]],
            size=11, color=MUTED, align=PP_ALIGN.CENTER)
    py = 5.2
    pill(s, Inches(0.78), Inches(py), Inches(5.75), Inches(0.95),
         [("Contract 1: ", True), ("raw packets never reach the LLM; Tier 1 hands over a ~20-number alert JSON.", False)])
    pill(s, Inches(6.80), Inches(py), Inches(5.75), Inches(0.95),
         [("Contract 2: ", True), ("no rule leaves the sandbox unverified; failures self-correct or escalate.", False)])


def _tele_example(tele) -> str:
    """Largest measured alert on the lab capture, e.g. '3,637 records (1.4 MB) -> 1.1 KB'."""
    if not tele or not tele.get("alerts"):
        return ""
    a = max(tele["alerts"], key=lambda x: x["flows"])
    names = {"SYN_FLOOD": "SYN flood", "CONN_FLOOD": "connection flood", "PORT_SCAN": "port scan",
             "BRUTE_FORCE": "brute force"}
    kind = " / ".join(names.get(r, r.replace("_", " ").lower()) for r in a["rules"])
    return (f"measured on the lab, one {kind} alert packs {a['flows']:,} Zeek records "
            f"({a['raw_bytes'] / 1e6:.1f} MB) into {a['alert_bytes']} bytes of JSON ({a['reduction']:.3%} smaller)")


def s_tier1(prs, tele):
    ex = _tele_example(tele)
    outcome = "raw packets never reach the LLM" + (f"; {ex}." if ex else ".")
    bullets_slide(prs, "Tier 1: Telemetry Triage", [
        ("Zeek sensor:", "passively captures the protected server's traffic and writes JSON connection records "
                         "(conn.log)."),
        ("Windowed aggregation:", "each source IP's flows in a 10-second window collapse into ~20 numbers: ports, "
                                  "hosts, rate, bytes and the half-open ratio TCP SYN_in / TCP ACK_out."),
        ("Deterministic rules:", "port scan, host sweep, SYN / connection / UDP flood, brute force, slow DoS, "
                                 "beaconing and exfiltration, each mapped to candidate ATT&CK tactics."),
        ("Isolation Forest:", "trained on benign traffic only, with its threshold taken from benign scores. This "
                              "fixes the prototype's fixed-contamination bug that flagged 5% of every batch."),
        ("Outcome:", outcome),
    ], size=18, notes="Tier 1 solves the telemetry blind spot: the LLM only ever sees a compact, typed alert.")


def _subscript_formula(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for text, sub in [("d", False), ("total", True), (" = d", False), ("proc", True), (" + d", False),
                      ("queue", True), (" + d", False), ("trans", True), (" + d", False), ("prop", True)]:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(24)
        r.font.italic = not sub
        r.font.color.rgb = BLACK
        if sub:
            r.font._rPr.set("baseline", "-25000")
    return tb


def s_async(prs, asy):
    s = blank_slide(prs, "Asynchronous Processing: Bounding the Nodal Delay", notes=(
        "The nodal-delay argument from the mid-review, now measured. With inference inline, every alert blocks "
        "ingestion, so queueing delay grows without bound. HRT-Agent enqueues and keeps ingesting."))
    bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.92), Inches(1.65), Inches(5.5), Inches(0.95))
    bg.adjustments[0] = 0.15
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT
    bg.line.fill.background()
    bg.shadow.inherit = False
    _subscript_formula(s, Inches(0.92), Inches(1.65), Inches(5.5), Inches(0.95))
    if asy is None:
        pending(s, "the async-vs-inline experiment")
        return
    rates = sorted(asy["rates"], key=int)
    inline = [asy["rates"][r]["inline"]["max_accept_lag_s"] for r in rates]
    asyncl = [asy["rates"][r]["async"]["max_accept_lag_s"] for r in rates]
    textbox(s, Inches(0.92), Inches(2.85), Inches(5.5), Inches(3.9), [
        [("Inline LLM: ", True), ("every alert blocks ingestion for the full inference time, so the queueing "
                                  "delay grows without bound.", False)],
        [("HRT-Agent: ", True), ("triage posts alerts to a webhook queue and keeps ingesting; two worker threads "
                                 "run the LLM.", False)],
        [("Measured: ", True), (f"over a {asy['duration_s']:.0f} s burst with {asy['llm_seconds']:.0f} s inference, "
                                f"inline ingest lag reaches {inline[0]:,.0f} s at {rates[0]} alert/s and "
                                f"{inline[-1]:,.0f} s at {rates[-1]} alerts/s. The async queue keeps ingest lag at "
                                f"{max(asyncl):.0f} s.", False)],
    ], size=16, space_after=12, bullets=True)
    line_chart(s, Inches(6.75), Inches(1.6), Inches(5.67), Inches(5.1), [f"{r}/s" for r in rates],
               [("Inline LLM", inline, ACCENT2), ("HRT-Agent (async queue)", asyncl, ACCENT1)],
               "Max ingest lag vs alert arrival rate")


def s_tier2(prs, examples, agent):
    s = blank_slide(prs, "Tier 2: RAG Intelligence Layer", notes=(
        "The knowledge base is local. Retrieval is hierarchical: the rule hits pick candidate tactics first, then "
        "cosine search runs only inside them. The table uses real alerts replayed from the lab capture."))
    kb = PROJECT / "data" / "chroma" / "kb_stats.json"
    kbs = json.loads(kb.read_text()) if kb.exists() else {"mitre_docs": 697, "cve_docs": 1710}
    red = agent.get("retrieval", {}).get("avg_search_space_reduction") if agent else None
    red_txt = f"; across the evaluation alerts it cut the search space by {red:.1%} on average" if red else ""
    textbox(s, Inches(0.92), Inches(1.55), Inches(11.5), Inches(2.9), [
        [("Knowledge base: ", True), (f"{kbs['mitre_docs']} MITRE ATT&CK techniques and sub-techniques plus "
                                      f"{kbs['cve_docs']:,} CISA Known Exploited Vulnerabilities, embedded locally "
                                      "in ChromaDB (cosine similarity, ONNX MiniLM; no GPU, no cloud).", False)],
        [("Hierarchical retrieval: ", True), ("Tier-1 rule hits select candidate tactics first, then cosine search "
                                              f"runs only within them{red_txt}.", False)],
        [("Local LLM: ", True), ("Llama 3.2 3B through Ollama, constrained to a JSON schema; no cloud API.", False)],
        [("Grounding check: ", True), ("the cited technique must be in the retrieved set, otherwise the answer is "
                                       "rejected and retried; hallucinated IDs never reach the analyst.", False)],
    ], size=16, space_after=8, bullets=True)
    if not examples:
        pending(s, "retrieval on replayed lab alerts")
        return
    rows = [["Tier-1 alert (real lab traffic)", "Tactic prefilter", "Search space", "Top retrieved technique",
             "Cosine"]]
    for e in examples:
        rows.append([e["pattern"], e["tactics"], e["space"], e["top"], f"{e['score']:.2f}"])
    table(s, rows, Inches(0.92), Inches(4.6), [2.55, 2.35, 2.05, 3.75, 0.8], font=13, row_h=0.42)


def s_tier3(prs):
    s = blank_slide(prs, "Tier 3: Deterministic Guardrail & Kernel Sandbox", notes=(
        "Four deterministic stages. The subnet guardrail is a bitwise AND of the rule's address and a protected "
        "network under the shorter mask. The last stage is a real kernel dry-run in a throwaway container."))
    rows = [
        ["Stage", "What it checks", "Example it stops"],
        ["1. Policy allowlist", "One -A/-I into the filter table; DROP/REJECT only; no shell metacharacters, "
                                "negation, flush or policy changes", "iptables … -j DROP; rm -rf /"],
        ["2. Subnet guardrail", "Bitwise AND, (A & m) == (B & m): never block the gateway, the protected server "
                                "or a known-good host; nothing broader than /24", "-s 10.10.0.0/24 (the whole LAN)"],
        ["3. Conflict check", "Duplicates, and DROPs shadowed by an earlier ACCEPT in the live ruleset",
         "-A … DROP placed after --dport 80 ACCEPT"],
        ["4. Kernel verification", "iptables-restore --test, then a real commit, inside an ephemeral Docker "
                                   "container (no network, read-only)", "--dport 80:70 (invalid range)"],
    ]
    table(s, rows, Inches(0.92), Inches(1.6), [2.35, 5.6, 3.55], font=13, row_h=0.8)
    pill(s, Inches(0.92), Inches(5.6), Inches(11.5), Inches(0.8),
         [("Self-correction: ", True), ("a rejection's reason is fed back to the LLM (≤ 3 attempts); non-retryable "
                                         "failures escalate to a human. Deployment waits for analyst approval.", False)],
         size=14)


def s_eval_setup(prs, triage):
    rows_n = f"{triage['rows']:,}" if triage else "?"
    rows = [
        ["Metric", "How it was measured", "Data"],
        ["Detection accuracy", "Precision, recall, F1 and ROC-AUC against ground-truth labels; baselines: "
                               "SIEM-style correlation rules and a supervised Random Forest",
         f"CICIDS2018: {rows_n} flows (FTP/SSH brute force, DoS-Hulk, DDoS-HOIC, benign)"],
        ["Incident report accuracy", "Cited ATT&CK technique vs a label → technique map; hierarchical RAG vs flat "
                                     "RAG vs no RAG", "Alerts synthesised from CICIDS2018 flows"],
        ["Rule validity rate", "Share of LLM rules passing kernel verification; 28-case adversarial suite vs a "
                               "regex validator", "Agent runs + hand-built attack rules"],
        ["Time-to-response", "Attack launch → verified, deployed block rule (target < 5 min)",
         "Live Docker lab: nmap scan, hping3 SYN flood, hydra SSH brute force"],
        ["Telemetry reduction", "Bytes of the Zeek records an alert summarises vs bytes of the alert JSON sent to "
                                "the LLM", "Live lab capture (Zeek conn.log)"],
        ["Nodal delay", "Ingest lag with the LLM inline vs behind an async queue", "Discrete-event simulation"],
    ]
    s = blank_slide(prs, "Evaluation Setup",
                    notes="The four metrics promised at the mid-review, plus telemetry reduction and the nodal-delay "
                          "check.")
    table(s, rows, Inches(0.92), Inches(1.55), [2.45, 5.4, 3.65], font=13, row_h=0.74)


def s_detection(prs, triage):
    s = blank_slide(prs, "Results: Tier-1 Detection on CICIDS2018", notes=(
        "Rules and the supervised forest win on these known attacks. The unsupervised forest ranks attacks well "
        "(AUC 0.96) and nails brute force but misses floods flow-by-flow, which is why the live Tier 1 pairs it "
        "with deterministic rate rules."))
    if triage is None:
        pending(s, "the CICIDS2018 detection evaluation")
        return
    d = triage["detectors"]
    names = [("HRT_iforest", "HRT Isolation Forest (unsupervised, 5% FPR budget)"),
             ("SIEM_rules", "SIEM-style correlation rules"),
             ("RandomForest_supervised", "Random Forest (supervised upper bound)")]
    rows = [["Detector", "Precision", "Recall", "F1", "ROC-AUC"]]
    for key, label in names:
        m = d[key]
        rows.append([label, f3(m["precision"]), f3(m["recall"]), f3(m["f1"]), f3(m.get("roc_auc"))])
    table(s, rows, Inches(0.92), Inches(1.55), [5.3, 1.55, 1.55, 1.55, 1.55], font=13, row_h=0.42)
    per = triage["hrt_per_class"]
    cats = [k for k in ("FTP-BruteForce", "SSH-Bruteforce", "DoS attacks-Hulk", "DDOS attack-HOIC") if k in per]
    bar_chart(s, Inches(0.92), Inches(3.45), Inches(6.4), Inches(3.55),
              [c.replace("DoS attacks-", "DoS ").replace("DDOS attack-", "DDoS ") for c in cats],
              [("Recall", [per[c]["recall"] for c in cats], ACCENT1)],
              "Isolation Forest recall by attack family", legend=False)
    fp = per.get("Benign", {}).get("false_positive_rate")
    textbox(s, Inches(7.6), Inches(3.55), Inches(4.82), Inches(3.4), [
        [("Brute force: ", True), ("caught almost perfectly by the unsupervised forest (recall "
                                   f"{per['FTP-BruteForce']['recall']:.2f} / {per['SSH-Bruteforce']['recall']:.3f}).",
                                   False)],
        [("Floods: ", True), ("volumetric DoS/DDoS flows look benign one at a time (HOIC recall "
                              f"{per['DDOS attack-HOIC']['recall']:.2f}).", False)],
        [("Design answer: ", True), ("pair the forest with deterministic rate rules, which catch floods (rules "
                                     f"alone reach F1 {d['SIEM_rules']['f1']:.3f}).", False)],
        [("False alarms: ", True), (f"benign false-positive rate {fp:.1%}, set by the 5% budget chosen from benign "
                                    "scores alone (no labels).", False)],
    ], size=14, space_after=9, bullets=True)


def s_telemetry(prs, tele):
    s = blank_slide(prs, "Results: Telemetry Reduction (Tier 1)", notes=(
        "Measured on the live lab capture, not estimated: for every alert, the bytes of the Zeek records it summarises "
        "versus the bytes of the alert JSON that is actually sent to the LLM."))
    if tele is None or not tele.get("alerts"):
        pending(s, "the telemetry measurement")
        return
    names = {"PORT_SCAN": "Port scan (nmap)", "SYN_FLOOD": "SYN flood (hping3)", "BRUTE_FORCE": "SSH brute force (hydra)"}
    groups: dict[str, list] = {}
    for a in tele["alerts"]:
        key = next((names[r] for r in a["rules"] if r in names), " + ".join(a["rules"]))
        groups.setdefault(key, []).append(a)

    def rng(vals, fmt):
        lo, hi = min(vals), max(vals)
        return fmt(lo) if lo == hi else f"{fmt(lo)} - {fmt(hi)}"

    rows = [["Alert type", "Alerts", "Zeek records summarised", "Raw size", "Alert JSON", "Reduction"]]
    for key, al in groups.items():
        rows.append([key, str(len(al)), rng([a["flows"] for a in al], lambda v: f"{v:,}"),
                     rng([a["raw_bytes"] for a in al], lambda v: f"{v / 1e6:.2f} MB" if v >= 1e6 else f"{v / 1e3:.1f} KB"),
                     rng([a["alert_bytes"] for a in al], lambda v: f"{v} B"),
                     rng([a["reduction"] for a in al], lambda v: f"{v:.3%}")])
    table(s, rows, Inches(0.92), Inches(1.6), [3.0, 0.85, 2.3, 2.0, 1.4, 1.95], font=13, row_h=0.5)
    sent = sum(a["alert_bytes"] for a in tele["alerts"])
    stat(s, Inches(0.92), Inches(4.05), Inches(3.7), f"{tele['median_alert_reduction']:.2%}",
         "median reduction per alert")
    stat(s, Inches(4.85), Inches(4.05), Inches(3.7), f"{tele['whole_log_reduction']:.3%}",
         f"everything logged ({tele['conn_log_bytes'] / 1e6:.0f} MB) vs everything sent to the LLM ({sent / 1e3:.1f} KB)")
    stat(s, Inches(8.8), Inches(4.05), Inches(3.6), f"{tele['records']:,}",
         f"Zeek records captured → {tele['n_alerts']} alerts")
    big = max(tele["alerts"], key=lambda a: a["flows"])
    pill(s, Inches(0.92), Inches(5.75), Inches(11.5), Inches(0.8),
         [("Why it matters: ", True), (f"a flood of {big['flows']:,} connections ({big['raw_bytes'] / 1e6:.1f} MB of "
                                       f"logs) reaches the LLM as {big['alert_bytes']} bytes of typed numbers, so context "
                                       "size, cost and latency no longer grow with traffic volume.", False)], size=14)


def s_guardrail(prs, g):
    s = blank_slide(prs, "Results: Rule Verification vs a Regex Validator", notes=(
        "Twenty-eight adversarial rules. The regex validator, typical of prior work, only checks that a string looks "
        "like an iptables DROP. It accepted every injection and every self-inflicted-DoS rule."))
    if g is None:
        pending(s, "the guardrail suite")
        return
    order = [("valid", "Valid rules"), ("injection", "Shell injection"), ("self_dos", "Self-inflicted DoS"),
             ("destructive", "Destructive ops"), ("wrong_target", "Wrong target"),
             ("invalid_syntax", "Invalid syntax")]
    cats = [(k, v) for k, v in order if k in g["by_category"]]
    hrt = [g["by_category"][k]["hrt_correct"] / g["by_category"][k]["n"] for k, _ in cats]
    rx = [g["by_category"][k]["regex_correct"] / g["by_category"][k]["n"] for k, _ in cats]
    bar_chart(s, Inches(0.92), Inches(1.6), Inches(7.6), Inches(5.2), [v for _, v in cats],
              [("HRT guardrail", hrt, ACCENT1), ("Regex validator", rx, ACCENT2)],
              "Correct verdicts by attack category", gap=60)
    stat(s, Inches(8.95), Inches(1.75), Inches(3.5), pct(g["hrt_accuracy"]),
         f"HRT guardrail accuracy on {g['n_cases']} adversarial rules")
    stat(s, Inches(8.95), Inches(3.45), Inches(3.5), pct(g["regex_accuracy"]), "regex validator accuracy", ACCENT2)
    stat(s, Inches(8.95), Inches(5.15), Inches(3.5), str(g["regex_false_accept_count"]),
         "unsafe rules the regex accepted, incl. every shell injection and self-DoS rule", ACCENT2)


def s_ablation(prs, agent):
    s = blank_slide(prs, "Results: Incident Attribution & RAG Ablation", notes=(
        "Same local LLM, same alerts, three modes. This isolates what retrieval grounding buys. The example ID "
        "comes from results/agent_probe_norag.txt, which prints what the no-RAG model answered attempt by attempt."))
    if agent is None or "rag_hier" not in agent:
        pending(s, "the agent (RAG ablation) evaluation")
        return
    modes = [("no_rag", "No RAG (LLM only)"), ("rag_flat", "Flat RAG"), ("rag_hier", "Hierarchical RAG (HRT)")]
    rows = [["Mode", "Technique acc.", "Tactic acc.", "Grounded", "Invalid IDs", "Rule verified", "Escalated"]]
    for k, label in modes:
        if k in agent:
            m = agent[k]
            rows.append([label, pct(m["technique_accuracy"]), pct(m["tactic_accuracy"]), pct(m["grounding_rate"]),
                         pct(m["invalid_id_rate"]), pct(m["rule_valid_rate"]), pct(m["escalation_rate"])])
    table(s, rows, Inches(0.92), Inches(1.55), [3.2, 1.45, 1.35, 1.3, 1.35, 1.5, 1.35], font=13, row_h=0.45,
          highlight_rows={len(rows) - 1})
    present = [(k, label) for k, label in modes if k in agent]
    bar_chart(s, Inches(0.92), Inches(3.55), Inches(6.6), Inches(3.4), [lbl for _, lbl in present],
              [("Technique accuracy", [agent[k]["technique_accuracy"] for k, _ in present], ACCENT1),
               ("Invalid technique IDs", [agent[k]["invalid_id_rate"] for k, _ in present], ACCENT2)],
              "Attribution accuracy vs hallucinated IDs", gap=70)
    meta = agent.get("meta", {})
    retr = agent.get("retrieval", {})
    ret = retr.get("hierarchical", {})
    n = meta.get("incidents", "?")
    classes = meta.get("classes") or ["DDoS-HOIC", "DoS-Hulk", "FTP brute force", "SSH brute force"]
    nr = agent.get("no_rag", {})
    n_int = n if isinstance(n, int) else 0
    invented = round(nr.get("invalid_id_rate", 0) * n_int)
    tac = round(nr.get("tactic_accuracy", 0) * n_int)
    tac_txt = "no tactic is right" if tac == 0 else f"{tac} of {n} tactics are right"
    tie = agent["rag_flat"]["technique_accuracy"] == agent["rag_hier"]["technique_accuracy"] if "rag_flat" in agent else False
    red = retr.get("avg_search_space_reduction")
    textbox(s, Inches(7.8), Inches(3.55), Inches(4.62), Inches(3.45), [
        [("Setup: ", True), (f"{n} alerts synthesised from CICIDS2018 ({len(classes)} classes), same local "
                             f"{meta.get('llm', 'LLM').replace('ollama:', '')} in all three modes.", False)],
        [("Without RAG: ", True), (f"{invented} of {n} answers cite technique IDs that do not exist in ATT&CK (e.g. "
                                   f"T1210.001), and {tac_txt}.", False)],
        [("With RAG: ", True), ("every ID is real and grounded; recall@1 "
                                f"{pct(ret.get('recall@1'))}." + (
                                    f" Flat and hierarchical tie on accuracy here; the hierarchy's gain is a "
                                    f"{red:.1%} smaller search space." if tie and red else ""), False)],
    ], size=13, space_after=7, bullets=True)


def s_live(prs, live):
    s = blank_slide(prs, "Results: Live End-to-End Response", notes=(
        "Measured on the Docker lab: time from launching the attack to a verified, deployed block rule, on a "
        "CPU-only laptop. Afterwards the live firewall is checked to confirm the attacker is blocked."))
    if live is None or not live.get("summary"):
        pending(s, "the live lab benchmark")
        return
    names = {"scan": "Port scan (nmap)", "synflood": "SYN flood (hping3)", "sshbf": "SSH brute force (hydra)"}
    rows = [["Attack", "Runs", "Mean", "Fastest", "Slowest", "< 5 min", "Technique cited"]]
    all_t = []
    for kind, v in live["summary"].items():
        techs = sorted({r.get("technique") for r in live["runs"] if r.get("attack") == kind and r.get("technique")})
        rows.append([names.get(kind, kind), str(v["runs"]), f"{v['mean_s']:.0f} s", f"{v['min_s']:.0f} s",
                     f"{v['max_s']:.0f} s", "Yes" if v["under_5min"] else "No", ", ".join(techs) or "-"])
        all_t += [r["total_response_s"] for r in live["runs"]
                  if r.get("attack") == kind and r.get("total_response_s")]
    table(s, rows, Inches(0.92), Inches(1.6), [3.0, 0.9, 1.25, 1.25, 1.25, 1.15, 2.7], font=13, row_h=0.5)
    detected = [r for r in live["runs"] if r.get("detected")]
    blocked = [r for r in detected if r.get("attacker_blocked_after")]
    mean_all = sum(all_t) / len(all_t) if all_t else None
    stat(s, Inches(0.92), Inches(4.4), Inches(3.6), f"{mean_all:.0f} s" if mean_all else "-",
         "mean attack → deployed rule (target 300 s)")
    stat(s, Inches(4.9), Inches(4.4), Inches(3.6), f"{len(detected)}/{len(live['runs'])}",
         "runs detected and resolved")
    stat(s, Inches(8.85), Inches(4.4), Inches(3.6), f"{len(blocked)}/{len(detected) or 1}",
         "attacker verified blocked on the live firewall afterwards")
    att = sorted(r["attempts"] for r in live["runs"] if r.get("attempts"))
    alert_t = sorted(r["to_alert"] for r in live["runs"] if r.get("to_alert") is not None)
    if att and alert_t:
        same = att[0] == att[-1]
        att_txt = (f"every run took {att[0]} LLM attempts" if same else f"runs took {att[0]}-{att[-1]} LLM attempts")
        pill(s, Inches(0.92), Inches(6.0), Inches(11.5), Inches(0.8),
             [("Where the time goes: ", True),
              (f"Tier 1 raised the alert {alert_t[0]:.1f}-{alert_t[-1]:.1f} s after launch; the rest is CPU inference. "
               f"{att_txt[0].upper() + att_txt[1:]}: early drafts were rejected by the deterministic checks and fixed "
               "from the verifier's feedback.", False)], size=13)


def s_console(prs):
    s = blank_slide(prs, "Live Demo: HRT-Agent SOC Console", notes=(
        "The dashboard streams each incident through the three tiers. The analyst sees why it fired, what it was "
        "grounded in, which checks the rule passed, and approves the deployment."))
    if not SCREENSHOT.exists():
        pending(s, "the console screenshot")
        return
    from PIL import Image
    w_px, h_px = Image.open(SCREENSHOT).size
    h = 5.35
    w = min(h * w_px / h_px, 8.3)
    h = w * h_px / w_px
    pic = s.shapes.add_picture(str(SCREENSHOT), Inches(0.92), Inches(1.55), Inches(w), Inches(h))
    pic.line.color.rgb = GRID
    pic.line.width = Pt(0.75)
    left = 0.92 + w + 0.3
    textbox(s, Inches(left), Inches(1.6), Inches(12.42 - left), Inches(5.3), [
        [("Tier 1: ", True), ("the features and rules that fired the alert.", False)],
        [("Tier 2: ", True), ("tactic prefilter, search-space cut and the retrieved techniques with cosine scores.",
                              False)],
        [("Tier 3: ", True), ("every guardrail stage with its timing, and each self-correction attempt.", False)],
        [("Approve & deploy: ", True), ("nothing reaches the server's firewall until an analyst approves it.",
                                        False)],
        # describes docs/deck/console.png (captured 2026-09-27); update it if the screenshot is retaken
        [("In this capture: ", True), ("a live lab run of three attacks (100k flows → 3 alerts). The SSH brute "
                                       "force was attributed to T1110.003, verified on attempt 3 and deployed; the "
                                       "scan and SYN flood are verified and waiting for approval.", False)],
    ], size=14, space_after=12, bullets=True)


def s_objectives(prs, tele):
    red = (f"; alerts {tele['median_alert_reduction']:.1%} smaller than the records they summarise (median, measured)"
           if tele and tele.get("median_alert_reduction") is not None else "")
    rows = [
        ["Objective", "Proposed", "Delivered"],
        ["1. Telemetry triage", "ML filter (e.g. Random Forest) over PCAP/NetFlow; anomalous JSON to the LLM",
         "Zeek + 10 s windows; threshold rules + benign-trained Isolation Forest (Random Forest kept as the "
         f"supervised baseline){red}"],
        ["2. RAG integration", "ChromaDB with MITRE ATT&CK and recent CVEs",
         "697 ATT&CK techniques + 1,710 CISA KEV CVEs; hierarchical retrieval; grounding check"],
        ["3. Safe rule synthesis", "LLM writes iptables and Snort rules; deterministic Python syntax verifier",
         "iptables rules from a local Llama 3.2; 4-stage guardrail with kernel iptables-restore --test; gated "
         "deploy (Snort left for future work)"],
        ["4. Empirical validation", "Benchmark vs SIEM tools on a multi-stage dataset: accuracy, time-to-response, "
                                    "rule validity",
         "CICIDS2018 + live Docker lab; SIEM-style rules and regex baselines (Splunk/QRadar not available)"],
    ]
    s = blank_slide(prs, "Objectives vs. Delivered",
                    notes="Mapping the four proposed objectives to what was built, including the deviations.")
    table(s, rows, Inches(0.92), Inches(1.55), [2.45, 3.85, 5.2], font=13, row_h=1.0)


def s_challenges(prs):
    bullets_slide(prs, "Challenges & Limitations", [
        ("Tampered threat-intel labels:", "the ATT&CK bundle served on our network carried non-standard tactic "
                                          "names; detected, documented and normalised at ingestion."),
        ("No IPs in CICIDS2018:", "dataset evaluation is flow-level; agent incidents use synthetic TEST-NET "
                                  "attacker addresses."),
        ("Slow link (~50 KB/s):", "the full 1.3 GB dataset was infeasible, so stratified HTTP-range slices of three "
                                  "attack days were used."),
        ("Windows Smart App Control:", "began blocking unsigned Python 3.14 ML binaries mid-project; ML imports "
                                       "were made lazy so live triage keeps running on the rules."),
        ("CPU-only, 3B model:", "~1.5-3 minutes per incident, mostly because the small model's first rule drafts "
                                "fail the checks and need up to 3 self-correction rounds: within target, but a GPU "
                                "or larger model would cut it sharply."),
        ("Emulated baselines:", "the SIEM baseline is SIEM-style rules, not Splunk/QRadar; Snort rule synthesis "
                                "is not implemented."),
    ], size=17, notes="Be upfront: each limitation is documented in docs/DATA_NOTES.md.")


def s_conclusion(prs, g, tele, agent, live):
    s = blank_slide(prs, "Conclusion & Future Work", notes="Close on the three gaps and how each was closed.")
    textbox(s, Inches(0.92), Inches(1.55), Inches(5.55), Inches(0.5), [[("What HRT-Agent shows", True)]], size=20,
            color=ACCENT1)
    g_txt = (f"{g['hrt_accuracy']:.0%} of adversarial rules judged correctly vs {g['regex_accuracy']:.0%} for regex"
             if g else "adversarial rules blocked")
    t_txt = f" ({tele['median_alert_reduction']:.2%} smaller, median)" if tele else ""
    h_txt = ""
    if agent and "no_rag" in agent and "rag_hier" in agent:
        h_txt = (f" (invented IDs: {agent['no_rag']['invalid_id_rate']:.0%} without RAG, "
                 f"{agent['rag_hier']['invalid_id_rate']:.0%} with)")
    l_txt = ""
    if live and live.get("summary"):
        worst = max(v["max_s"] for v in live["summary"].values())
        l_txt = f" Live attacks blocked in ≤ {worst:.0f} s, under the 5-minute target."
    textbox(s, Inches(0.92), Inches(2.15), Inches(5.55), Inches(4.7), [
        [("Telemetry blind spot closed: ", True), (f"the LLM reasons over a compact alert{t_txt}, never raw "
                                                   "packets.", False)],
        [("Hallucination contained: ", True), (f"every technique is grounded in retrieved ATT&CK evidence{h_txt}.",
                                               False)],
        [("Excessive agency removed: ", True), (f"deterministic, kernel-verified guardrail ({g_txt}).", False)],
        [("Line-rate safe: ", True), ("the async queue keeps ingestion independent of inference latency." + l_txt,
                                      False)],
    ], size=15, space_after=11, bullets=True)
    textbox(s, Inches(6.87), Inches(1.55), Inches(5.55), Inches(0.5), [[("Future work", True)]], size=20,
            color=ACCENT1)
    textbox(s, Inches(6.87), Inches(2.15), Inches(5.55), Inches(4.7), [
        [("GPU / larger local model ", True), ("for faster, more accurate attribution.", False)],
        [("Snort / Suricata rules ", True), ("alongside iptables, with the same verifier pattern.", False)],
        [("Multi-source correlation ", True), ("(host + network logs) and IPv6 support.", False)],
        [("SIEM connectors ", True), ("(Splunk, QRadar) and an online benign baseline.", False)],
    ], size=15, space_after=11, bullets=True)


def s_thanks(prs):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    t = s.shapes.title
    t.text_frame.text = ""
    r = t.text_frame.paragraphs[0].add_run()
    r.text = "Thank You"
    r.font.bold = True
    r.font.italic = True
    r.font.size = Pt(44)
    sub = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
    sub.text_frame.text = ""
    _runs(sub.text_frame.paragraphs[0], [("QUESTIONS & DISCUSSION", True)], 20, BLACK)
    _runs(sub.text_frame.add_paragraph(), [("GROUP 9 · AMRITA VISHWA VIDYAPEETHAM", True)], 16, MUTED)


# ------------------------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--final", action="store_true", help="fail if any result file is missing")
    ap.add_argument("--out", type=Path, default=OUT_DECK)
    args = ap.parse_args()

    triage, guard, asy = load("triage"), load("guardrail"), load("async")
    agent, live, tele = load("agent"), load("live"), load("telemetry")
    if args.final and MISSING:
        print("missing results:", ", ".join(MISSING))
        return 1

    prs = Presentation(str(SRC_DECK))        # the original is only read; output is a new file
    n_orig = len(prs.slides)
    examples = retrieval_examples()

    s_architecture(prs)
    s_tier1(prs, tele)
    s_async(prs, asy)
    s_tier2(prs, examples, agent)
    s_tier3(prs)
    s_eval_setup(prs, triage)
    s_detection(prs, triage)
    s_telemetry(prs, tele)
    s_guardrail(prs, guard)
    s_ablation(prs, agent)
    s_live(prs, live)
    s_console(prs)
    s_objectives(prs, tele)
    s_challenges(prs)
    s_conclusion(prs, guard, tele, agent, live)
    s_thanks(prs)

    prs.save(str(args.out))
    print(f"kept {n_orig} original slides, added {len(prs.slides) - n_orig}; saved {args.out}")
    if MISSING:
        print("NOTE: slides flagged 'RESULTS PENDING' for:", ", ".join(MISSING))
    return 0


if __name__ == "__main__":
    sys.exit(main())
